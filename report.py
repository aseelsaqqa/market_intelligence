
import os
import pandas as pd

from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


def create_excel_report(
    review_summary_path,
    rating_distribution_path,
    sentiment_summary_path,
    output_path="data/competitor_intelligence_report.xlsx",
):
    """Create a formatted Excel report for competitor intelligence."""

    files = {
        "Competitor Overview": review_summary_path,
        "Rating Distribution": rating_distribution_path,
        "Sentiment Analysis": sentiment_summary_path,
    }

    for sheet_name, file_path in files.items():
        if not os.path.exists(file_path):
            raise FileNotFoundError(
                f"Required analysis file not found: {file_path}"
            )

    output_dir = os.path.dirname(output_path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    dataframes = {
        sheet_name: pd.read_csv(file_path)
        for sheet_name, file_path in files.items()
    }

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:

        # Write existing analysis sheets
        for sheet_name, df in dataframes.items():
            df.to_excel(writer, sheet_name=sheet_name, index=False)

        # Build executive summary from existing data
        overview = dataframes["Competitor Overview"]
        sentiment = dataframes["Sentiment Analysis"]

        summary_rows = []

        for _, row in overview.iterrows():
            title = row.get("title", "Unknown")
            app_id = row.get("app_id", "")
            rating = row.get("average_rating", None)
            review_count = row.get("review_count", None)

            sentiment_row = sentiment[
                sentiment["app_id"].astype(str) == str(app_id)
            ] if "app_id" in sentiment.columns else pd.DataFrame()

            positive_pct = None
            negative_pct = None

            if not sentiment_row.empty:
                s = sentiment_row.iloc[0]
                positive_pct = s.get("Positive %", None)
                negative_pct = s.get("Negative %", None)

            summary_rows.append({
                "Competitor": title,
                "Average Rating": rating,
                "Review Count": review_count,
                "Positive Reviews (%)": positive_pct,
                "Negative Reviews (%)": negative_pct,
            })

        pd.DataFrame(summary_rows).to_excel(
            writer, sheet_name="Executive Summary", index=False
        )

        # Format all worksheets
        workbook = writer.book

        for ws in workbook.worksheets:
            ws.freeze_panes = "A2"
            ws.sheet_view.showGridLines = False
            ws.auto_filter.ref = ws.dimensions

            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(
                    "solid", fgColor="1F4E78"
                )
                cell.alignment = Alignment(
                    horizontal="center", vertical="center"
                )

            ws.row_dimensions[1].height = 30

            for column_cells in ws.columns:
                max_length = max(
                    len(str(cell.value or "")) for cell in column_cells
                )
                column_letter = get_column_letter(
                    column_cells[0].column
                )
                ws.column_dimensions[column_letter].width = min(
                    max(max_length + 3, 14), 35
                )

            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4

        # Add a chart to the rating distribution sheet
        ws = workbook["Rating Distribution"]

        if ws.max_row > 1 and ws.max_column > 1:
            chart = BarChart()
            chart.type = "col"
            chart.style = 10
            chart.title = "Rating Distribution by Competitor"
            chart.y_axis.title = "Number of Reviews"
            chart.x_axis.title = "Competitor"
            chart.height = 8
            chart.width = 16

            data = Reference(
                ws, min_col=2, max_col=ws.max_column,
                min_row=1, max_row=ws.max_row
            )
            categories = Reference(
                ws, min_col=1, min_row=2, max_row=ws.max_row
            )

            chart.add_data(data, titles_from_data=True)
            chart.set_categories(categories)
            ws.add_chart(chart, f"A{ws.max_row + 3}")

    return output_path
